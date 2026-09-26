from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from .field_extractor import extract_candidates
from .models import ExtractedCandidate, ExtractionResult


def _workbook_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows: list[str] = []
    for worksheet in workbook.worksheets:
        rows.append(f"Sheet: {worksheet.title}")
        for row in worksheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(": ".join(values) if len(values) == 2 else " | ".join(values))
    return "\n".join(rows)


def process_spreadsheet(content: bytes) -> ExtractionResult:
    text = _workbook_text(content)
    if not text.strip():
        raise ValueError("EMPTY_SPREADSHEET")
    candidates = extract_candidates([text], source="SPREADSHEET")
    return ExtractionResult(text, 1, candidates, "form_16" if any(item.field in {"salary_income", "tds"} for item in candidates) else "other")
